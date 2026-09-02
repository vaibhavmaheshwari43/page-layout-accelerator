#!/usr/bin/env python3
"""
Veeva -> LSC Page Layout Accelerator — Decision Matrix Report Generator

This project started from a manually-built Excel Decision Matrix (the
original SP_Admin_Layout_HCO_Migration_Decision_Matrix.xlsx). This script
closes the loop: given ANY completed pipeline run (the classified_*.json
output from rules_engine.py), it automatically produces that same style of
report -- for stakeholder review and sign-off -- instead of someone building
it by hand each time.

Works for any layout, any object -- not tied to HCO. Point it at HCP's
classified output, or any future layout's, and it produces the same
professional report format.

Usage:
    python3 decision_matrix_report.py --classified classified_SP_Admin_Layout_HCO.json \
        --out SP_Admin_Layout_HCO_Decision_Matrix_Report.xlsx
"""

import argparse
import json
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")

ACTION_FILL = {
    "auto_generate": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "flag_manual_review": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "flag_no_registry_entry": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "flag_rebuild": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "flag_retire": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "flag_decision_needed": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "informational_only": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}

ACTION_LABELS = {
    "auto_generate": "Auto-build",
    "flag_manual_review": "Needs quick confirmation",
    "flag_no_registry_entry": "No mapping yet -- needs review",
    "flag_rebuild": "Needs custom rebuild",
    "flag_retire": "Proposed for retirement",
    "flag_decision_needed": "Needs a business decision",
    "informational_only": "Informational only",
}


def build_report(classified_path, out_path):
    with open(classified_path, encoding="utf-8") as f:
        data = json.load(f)

    elements = [e for e in data["elements"] if not e["element_type"].startswith("_")]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Decision Matrix"

    headers = ["Layout Section", "Element Type", "API Name", "Classification",
               "LSC Target", "Confidence", "Status", "Org Verified", "Basis"]

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, e in enumerate(elements, start=2):
        target = e.get("target", {})
        row = [
            e["layout_section"],
            e["element_type"],
            e["api_name"],
            e.get("classification", {}).get("canonical", "-"),
            target.get("api_name") or "-",
            target.get("confidence", "-"),
            ACTION_LABELS.get(e.get("action", ""), e.get("action", "-")),
            ("Yes" if e.get("org_verified") else ("No" if "org_verified" in e else "-")),
            e.get("basis", "-"),
        ]
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = ACTION_FILL.get(e.get("action"))
        if fill:
            ws.cell(row=i, column=7).fill = fill

    widths = [22, 18, 34, 16, 30, 12, 24, 12, 50]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(elements) + 1}"

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = f"Decision Matrix Report -- {data.get('source_object', '?')} / {data.get('source_layout', '?')}"
    ws2["A1"].font = Font(bold=True, size=13)

    action_counts = Counter(e.get("action") for e in elements)
    ws2["A3"] = "Status"
    ws2["B3"] = "Count"
    ws2["A3"].font = Font(bold=True)
    ws2["B3"].font = Font(bold=True)
    row = 4
    for action, label in ACTION_LABELS.items():
        count = action_counts.get(action, 0)
        if count:
            ws2.cell(row=row, column=1, value=label)
            ws2.cell(row=row, column=2, value=count)
            fill = ACTION_FILL.get(action)
            if fill:
                ws2.cell(row=row, column=1).fill = fill
                ws2.cell(row=row, column=2).fill = fill
            row += 1

    row += 1
    ws2.cell(row=row, column=1, value="Total elements").font = Font(bold=True)
    ws2.cell(row=row, column=2, value=len(elements)).font = Font(bold=True)
    row += 1
    matched = sum(1 for e in elements if e.get("registry_match"))
    ws2.cell(row=row, column=1, value="Registry-matched")
    ws2.cell(row=row, column=2, value=matched)

    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 12

    wb.save(out_path)
    return {"total_elements": len(elements), "action_counts": dict(action_counts)}


def main():
    ap = argparse.ArgumentParser(description="Generate a stakeholder-ready Decision Matrix report from a classified pipeline run.")
    ap.add_argument("--classified", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    result = build_report(args.classified, args.out)
    print(f"Report written to {args.out}")
    print(f"Total elements: {result['total_elements']}")
    print(f"Action breakdown: {result['action_counts']}")


if __name__ == "__main__":
    main()
